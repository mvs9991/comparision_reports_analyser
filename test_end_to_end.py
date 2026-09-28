"""End-to-end checks: whole runs through recon_html_report.py, workbook and page compared.

test_patterns.py proves each value-pair rule on its own. This file proves the rest
of the promise: a broken input is reported and never shown as clean, the workbook
and the HTML page carry the same numbers, and nothing the page shows can break it.
Every run happens in a temporary folder that is removed afterwards.

Run:  python test_end_to_end.py      (exit code 0 = all pass)
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from openpyxl import load_workbook  # noqa: E402

import recon_analyzer as ra  # noqa: E402

PY = sys.executable
FAILS = []


def check(name, ok, detail=""):
    print(f"  {'ok ' if ok else 'FAIL'} {name}" + ("" if ok else f"   ({detail})"))
    if not ok:
        FAILS.append(name)


def write_csv(folder, table, lines):
    path = os.path.join(folder, table)
    os.makedirs(path, exist_ok=True)
    with open(os.path.join(path, f"{table}_mismatch.csv"), "w", encoding="utf-8", newline="") as fh:
        fh.write("\n".join(lines) + "\n")


def run_report(run_folder, out_dir, *extra):
    """recon_html_report.py on a run folder; returns (exit code, output, xlsx, html)."""
    run = os.path.basename(os.path.normpath(run_folder))
    xlsx = os.path.join(out_dir, f"report_summary_{run}.xlsx")
    proc = subprocess.run(
        [PY, os.path.join(HERE, "recon_html_report.py"), run_folder, "-o", xlsx, *extra],
        capture_output=True, text=True, encoding="utf-8", errors="replace",
        env=dict(os.environ, PYTHONIOENCODING="utf-8"),
    )
    return proc.returncode, proc.stdout + proc.stderr, xlsx, os.path.splitext(xlsx)[0] + ".html"


def page_data(html_path):
    with open(html_path, encoding="utf-8") as fh:
        text = fh.read()
    m = re.search(r'<script id="data" type="application/json">(.*?)</script>', text, re.S)
    return json.loads(m.group(1)), m.group(1), text


def workbook_columns(xlsx_path, sheet):
    """{column: {pattern: n}} from a table sheet's column-level summary."""
    ws = load_workbook(xlsx_path)[sheet]
    out, in_block = {}, False
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "Risk" and row[1] == "Column":
            in_block = True
            continue
        if in_block:
            if not row or row[0] is None:
                break
            out[row[1]] = {p: int(n) for p, n in (part.split("=") for part in row[5].split(", "))}
    return out


def test_sample_run(tmp):
    print("\n--- sample_input: clean run, workbook == page ---")
    code, out, xlsx, html = run_report(os.path.join(HERE, "sample_input", "20260907"), tmp)
    check("exits 0", code == 0, out[-300:])
    wb = load_workbook(xlsx)
    errors = [r for r in wb["Errors"].iter_rows(min_row=2, values_only=True)]
    check("Errors sheet says no problems", errors and errors[0][:2] == ("(none)", "no problems"), errors)
    view, _, _ = page_data(html)
    check("page lists both tables", sorted(t["table"] for t in view["tables"])
          == ["customer_acct_recon", "demo_small_recon"])
    for t in view["tables"]:
        page_cols = {c["column"]: {p["name"]: p["n"] for p in c["patterns"]} for c in t["columns"]}
        check(f"{t['table']}: same columns and pattern counts in workbook and page",
              page_cols == workbook_columns(xlsx, t["table"]),
              f"page {page_cols} vs workbook {workbook_columns(xlsx, t['table'])}")


def test_broken_inputs(tmp):
    print("\n--- broken inputs are reported, never shown as clean ---")
    run = os.path.join(tmp, "broken_run")
    os.makedirs(run)
    with open(os.path.join(run, "stray.csv"), "w") as fh:
        fh.write("a,b\n")
    os.makedirs(os.path.join(run, "no_csv_recon"))
    write_csv(run, "bad_header_recon", ["foo,bar", "1,2"])
    write_csv(run, "empty_recon", [])
    write_csv(run, "messy_recon", [
        'match_type,mismatch_columns,id,amt__s2,amt__s3,note__s2,note__s3',
        'mismatch,amt,1,10,11,a,a',
        'mismatch,amt,1,10,12,a,a',               # same key twice
        'mismatch,amt,2,10,10,a,b',               # note differs but was not listed
        'mismatch,ghost,3,1,1,a,a',               # names a column the file does not have
        'added,amt,4,1,2,a,a',                    # not a mismatch row
        'mismatch,amt,5,1',                       # too few fields
        'mismatch,,6,1,1,a,a',                    # blank list, nothing differs
    ])
    code, out, xlsx, html = run_report(run, tmp, "--summary")
    check("exits 0 despite broken tables", code == 0, out[-300:])
    view, _, _ = page_data(html)
    not_analysed = {e["table"] for e in view["errors"]}
    for name in ("bad_header_recon", "empty_recon", "no_csv_recon", "(files directly in the input folder)"):
        check(f"NOT ANALYSED: {name}", name in not_analysed, sorted(not_analysed))
    messy = next(t for t in view["tables"] if t["table"] == "messy_recon")
    notes = " | ".join(messy["notes"])
    for needle in ("DUPLICATE PRIMARY KEYS", "RECON DID NOT FLAG: 'note'", "ghost",
                   "match_type did not say mismatch", "fewer fields than the header",
                   "identical s2/s3 text in every column"):
        check(f"warning mentions {needle!r}", needle in notes, notes)
    check("messy table needs a manual check", bool(messy["reasons"]))
    errors_sheet = [r[1] for r in load_workbook(xlsx)["Errors"].iter_rows(min_row=2, values_only=True)]
    check("Errors sheet lists 4 NOT ANALYSED rows", errors_sheet.count("NOT ANALYSED") == 4, errors_sheet)


def test_one_sided_column(tmp):
    print("\n--- a column on one side only is reported ---")
    run = os.path.join(tmp, "one_sided")
    write_csv(run, "t_recon", [
        'match_type,mismatch_columns,id,amt__s2,amt__s3,legacy__s2,newcol__s3',
        'mismatch,amt,1,10,11,X,Y',
    ])
    r = ra.analyze_table("t_recon", os.path.join(run, "t_recon", "t_recon_mismatch.csv"))
    notes = " | ".join(r["notes"])
    check("legacy__s2 reported as s2-only", "legacy__s2" in notes and "only as __s2" in notes, notes)
    check("newcol__s3 reported as s3-only", "newcol__s3" in notes and "only as __s3" in notes, notes)
    check("raises a manual check", any("one side only" in x for x in r["manual_check_reasons"]),
          r["manual_check_reasons"])


def test_match_type_not_a_key(tmp):
    print("\n--- match_type after mismatch_columns is not part of the key ---")
    run = os.path.join(tmp, "mt_after")
    write_csv(run, "t_recon", [
        'mismatch_columns,match_type,id,amt__s2,amt__s3',
        'amt,mismatch,1,10,11',
    ])
    r = ra.analyze_table("t_recon", os.path.join(run, "t_recon", "t_recon_mismatch.csv"))
    check("primary key is just id", r["pk_cols"] == ["id"], r["pk_cols"])


def test_column_overrides(tmp):
    print("\n--- free-text and yes/no columns are read by their values ---")
    run = os.path.join(tmp, "overrides")
    write_csv(run, "t_recon", [
        'match_type,mismatch_columns,id,remarks__s2,remarks__s3,active__s2,active__s3',
        'mismatch,remarks,1,by post,by mail,Y,Y',
        'mismatch,active,2,x,x,1,true',
        'mismatch,active,3,x,x,N,0',
    ])
    r = ra.analyze_table("t_recon", os.path.join(run, "t_recon", "t_recon_mismatch.csv"))
    cols = {c["column"]: c["patterns"] for c in r["columns"]}
    check("remarks rewording -> text_field_diff", cols.get("remarks") == {"text_field_diff": 1}, cols)
    check("yes/no column: 1 vs true -> boolean_format_diff",
          cols.get("active", {}).get("boolean_format_diff") == 2, cols)


def test_bad_template_fails_fast(tmp):
    print("\n--- a bad --recon-table stops before any analysis ---")
    code, out, xlsx, _ = run_report(os.path.join(HERE, "sample_input", "20260907"),
                                    os.path.join(tmp, "bad_tpl"), "--recon-table", "{schema}.{table}")
    check("exits non-zero", code != 0)
    check("says why", "--recon-table" in out and "{table}" in out, out[-300:])
    check("writes no workbook", not os.path.exists(xlsx))


def test_output_folder(tmp):
    print("\n--- -o <existing folder> writes inside it ---")
    out_dir = os.path.join(tmp, "out_folder")
    os.makedirs(out_dir)
    proc = subprocess.run([PY, os.path.join(HERE, "recon_html_report.py"),
                           os.path.join(HERE, "sample_input", "20260907"), "-o", out_dir],
                          capture_output=True, text=True)
    check("exits 0", proc.returncode == 0, proc.stderr[-300:])
    for ext in ("xlsx", "html"):
        check(f"report_summary_20260907.{ext} inside the folder",
              os.path.isfile(os.path.join(out_dir, f"report_summary_20260907.{ext}")))


def test_hostile_values_cannot_break_page(tmp):
    print("\n--- exported values cannot break the page ---")
    run = os.path.join(tmp, "hostile")
    write_csv(run, "t_recon", [
        'match_type,mismatch_columns,id,note__s2,note__s3',
        'mismatch,note,1,"<!--<script>alert(1)</script>","x</script><b>y"',
    ])
    code, out, _, html = run_report(run, tmp)
    check("exits 0", code == 0, out[-300:])
    view, raw, text = page_data(html)
    check("no raw '<' inside the data block", "<" not in raw)
    check("exactly two </script> tags (data + code)", text.count("</script>") == 2, text.count("</script>"))
    ex = view["tables"][0]["columns"][0]["examples"][0]
    check("values survive exactly", ex["s2"] == "<!--<script>alert(1)</script>" and ex["s3"] == "x</script><b>y", ex)


def main():
    tmp = tempfile.mkdtemp(prefix="recon_e2e_")
    try:
        for test in (test_sample_run, test_broken_inputs, test_one_sided_column,
                     test_match_type_not_a_key, test_column_overrides,
                     test_bad_template_fails_fast, test_output_folder,
                     test_hostile_values_cannot_break_page):
            try:
                test(tmp)
            except Exception as exc:  # a crash is a failure, not a stop
                check(f"{test.__name__} ran without crashing", False, f"{type(exc).__name__}: {exc}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("\n" + "=" * 78)
    print(f"{len(FAILS)} end-to-end check(s) FAILED: {FAILS}" if FAILS else "all end-to-end checks passed")
    return 1 if FAILS else 0


if __name__ == "__main__":
    sys.exit(main())
